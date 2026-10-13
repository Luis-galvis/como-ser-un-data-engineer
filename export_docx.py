import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

CAP_DIR = r'c:\Users\lgalv\Desktop\obsidian\libro\capitulos'
OUT_DOCX = r'c:\Users\lgalv\Desktop\obsidian\libro\libro-profesional.docx'
MIRROR_DOCX = r'D:\estudio\libro\libro-profesional.docx'

print("Iniciando compilacion de documento DOCX academico...")

doc = Document()

# Configurar tamano de pagina 7x10 in (Handicraft Academic Book)
for section in doc.sections:
    section.page_width = Inches(7.0)
    section.page_height = Inches(10.0)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.75)

# 1. Portada DOCX
p_pre = doc.add_paragraph()
p_pre.alignment = WD_ALIGN_PARAGRAPH.LEFT
r_pre = p_pre.add_run("Cómo ser un")
r_pre.font.name = "Georgia"
r_pre.font.size = Pt(20)
r_pre.font.italic = True
r_pre.font.color.rgb = RGBColor(74, 81, 96)

p_title = doc.add_paragraph()
r_title = p_title.add_run("DATA ENGINEER")
r_title.font.name = "Georgia"
r_title.font.size = Pt(36)
r_title.font.bold = True
r_title.font.color.rgb = RGBColor(27, 32, 48)

p_sub = doc.add_paragraph()
r_sub = p_sub.add_run("Una mirada a fondo de los fundamentos y los procesos prácticos para llegar de cero a Data Engineer")
r_sub.font.name = "Georgia"
r_sub.font.size = Pt(13)
r_sub.font.italic = True
r_sub.font.color.rgb = RGBColor(74, 81, 96)

# Espaciado e imagen
condor_path = r'c:\Users\lgalv\Desktop\obsidian\condor_profesional.jpg'
if os.path.exists(condor_path):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(condor_path, width=Inches(3.8))

p_author = doc.add_paragraph()
p_author.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r_author = p_author.add_run("\n\nLuis Galvis\n2026")
r_author.font.name = "Georgia"
r_author.font.size = Pt(14)
r_author.font.bold = True
r_author.font.color.rgb = RGBColor(27, 32, 48)

doc.add_page_break()

# 2. Pagina Legal
p_leg_title = doc.add_paragraph()
r_lt = p_leg_title.add_run("Cómo ser un Data Engineer\n")
r_lt.font.bold = True
r_lt.font.size = Pt(12)
p_leg_title.add_run("Primera edición: Octubre de 2026\nAutor: Luis Galvis\n© 2026 Luis Galvis. Todos los derechos reservados.\n\n")

p_leg_body = doc.add_paragraph()
p_leg_body.add_run(
    "Ninguna parte de esta obra puede ser reproducida, almacenada en sistemas de recuperación ni transmitida, "
    "total o parcialmente, por ningún medio, sin la autorización previa y por escrito del autor.\n\n"
    "Los fragmentos de código fuente (PySpark, SQL, Python, HCL) se distribuyen bajo la licencia de código abierto MIT "
    "y pueden reutilizarse libremente en entornos productivos, conservando el aviso de licencia.\n\n"
    "Escrito y compilado en Ibagué, Tolima, Colombia • 2026"
)

doc.add_page_break()

# 3. Dedicatoria
p_ded = doc.add_paragraph()
p_ded.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_ded = p_ded.add_run(
    "\n\n\n\n\n\n\"A mis padres, por ser el cimiento inquebrantable, el sacrificio silencioso y la luz que impulsó cada uno de mis pasos. "
    "A mi pareja, por su amor incondicional, su infinita paciencia y su aliento constante en cada noche de estudio y creación. "
    "Y a todos los ingenieros y científicos que buscan comprender la esencia profunda de los datos más allá de las herramientas efímeras.\"\n\n— Luis Galvis"
)
r_ded.font.name = "Georgia"
r_ded.font.italic = True
r_ded.font.size = Pt(13)
r_ded.font.color.rgb = RGBColor(74, 81, 96)

doc.add_page_break()

# 4. Capítulos
files = [f for f in sorted(os.listdir(CAP_DIR)) if f.endswith('.md') and f != '00-portada.md']

def sanitize_xml(s):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', s)

for file in files:
    filepath = os.path.join(CAP_DIR, file)
    with open(filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code = False
    code_lines = []

    for line in lines:
        stripped = line.strip()

        # Codigo
        if stripped.startswith('```'):
            if in_code:
                in_code = False
                code_text = sanitize_xml("".join(code_lines))
                p_code = doc.add_paragraph()
                p_code.paragraph_format.left_indent = Inches(0.3)
                p_code.paragraph_format.right_indent = Inches(0.3)
                r_c = p_code.add_run(code_text)
                r_c.font.name = "Consolas"
                r_c.font.size = Pt(8.5)
                r_c.font.color.rgb = RGBColor(27, 32, 48)
                code_lines = []
            else:
                in_code = True
                code_lines = []
            continue

        if in_code:
            code_lines.append(line)
            continue

        # Encabezados
        if stripped.startswith('# '):
            h1_text = sanitize_xml(stripped[2:].strip())
            h1 = doc.add_heading(h1_text, level=1)
            h1.paragraph_format.space_before = Pt(18)
            h1.paragraph_format.space_after = Pt(8)
            for r in h1.runs:
                r.font.name = "Georgia"
                r.font.color.rgb = RGBColor(27, 32, 48)
        elif stripped.startswith('## '):
            h2_text = sanitize_xml(stripped[3:].strip())
            h2 = doc.add_heading(h2_text, level=2)
            h2.paragraph_format.space_before = Pt(14)
            h2.paragraph_format.space_after = Pt(6)
            for r in h2.runs:
                r.font.name = "Georgia"
                r.font.color.rgb = RGBColor(27, 32, 48)
        elif stripped.startswith('### '):
            h3_text = sanitize_xml(stripped[4:].strip())
            h3 = doc.add_heading(h3_text, level=3)
            h3.paragraph_format.space_before = Pt(10)
            h3.paragraph_format.space_after = Pt(4)
            for r in h3.runs:
                r.font.name = "Georgia"
                r.font.color.rgb = RGBColor(74, 81, 96)
        elif stripped.startswith('> '):
            quote_text = sanitize_xml(stripped[2:].strip().replace('**', '').replace('*', ''))
            p_q = doc.add_paragraph()
            p_q.paragraph_format.left_indent = Inches(0.4)
            p_q.paragraph_format.space_before = Pt(4)
            p_q.paragraph_format.space_after = Pt(4)
            r_q = p_q.add_run(quote_text)
            r_q.font.name = "Georgia"
            r_q.font.italic = True
            r_q.font.size = Pt(9.5)
            r_q.font.color.rgb = RGBColor(74, 81, 96)
        elif stripped.startswith('* ') or stripped.startswith('- '):
            list_text = sanitize_xml(stripped[2:].strip().replace('**', '').replace('*', ''))
            p_l = doc.add_paragraph(list_text, style='List Bullet')
            p_l.paragraph_format.space_after = Pt(2)
        elif stripped.startswith('---'):
            continue
        elif stripped:
            clean_text = sanitize_xml(stripped.replace('**', '').replace('*', ''))
            p_b = doc.add_paragraph(clean_text)
            p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_b.paragraph_format.first_line_indent = Inches(0.25)
            p_b.paragraph_format.space_after = Pt(4)

    doc.add_page_break()


# Guardar DOCX
doc.save(OUT_DOCX)
print(f"[OK] Documento DOCX generado: {OUT_DOCX}")

if os.path.exists(r'D:\estudio\libro'):
    doc.save(MIRROR_DOCX)
    print(f"[OK] Documento DOCX sincronizado en: {MIRROR_DOCX}")

